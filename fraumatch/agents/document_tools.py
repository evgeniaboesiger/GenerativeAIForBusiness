"""
MATCHA - Document processing tools for the Profile Agent.

Standard PDF text extraction (``pypdf``) is fast but can fail or return
incomplete/incorrect text - especially for scanned (image-based) PDFs where
the text is drawn as pixels instead of a real text layer.

This module is a supporting TOOL used by the Profile Agent (NOT a separate
AI agent): it converts raw document bytes into plain text. OCR is only used
as a *fallback* when the fast extraction is clearly insufficient, so it
adds processing time only when it can actually help.

Extraction result sources:
- ``txt``, ``docx``            fast text extraction
- ``pdf_text``                 fast PDF text layer extraction
- ``pdf_ocr``                  OCR fallback used (image-based / broken PDF)
- ``pdf_empty``                PDF produced no usable text and OCR was skipped
- ``pdf_ocr_missing``          PDF needs OCR but no OCR engine is installed
- ``pdf_pkg_missing``          pypdf not installed
- ``docx_pkg_missing``         python-docx not installed
- ``unsupported``              unknown file type
"""

import io
import re
from dataclasses import dataclass

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover - environment missing dependency
    PdfReader = None

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover - environment missing dependency
    DocxDocument = None

# Minimum content a CV should have before we treat fast text extraction as
# sufficient. A real CV - even a short one - has far more than this.
_MIN_WORDS = 25
_MIN_CHARS = 150


@dataclass
class CVTextResult:
    """Plain text produced from a document plus the method used."""

    text: str
    source: str = "unsupported"

    @property
    def chars(self) -> int:
        return len(self.text)

    @property
    def ocr_used(self) -> bool:
        return self.source == "pdf_ocr"


def is_insufficient(text: str) -> bool:
    """Return True when fast text extraction clearly produced too little.

    Scanned / image-only PDFs typically yield an empty string or a handful
    of stray characters from ``pypdf``, which triggers the OCR fallback.
    """
    if not text:
        return True
    words = re.findall(r"\b\w+\b", text)
    chars = len(re.sub(r"\s+", "", text))
    return len(words) < _MIN_WORDS or chars < _MIN_CHARS


class OCREngine:
    """OCR engine: renders PDF pages with PyMuPDF and reads them with Tesseract.

    Everything is imported lazily so the application keeps working even when
    the optional OCR dependencies (or the Tesseract binary) are missing.
    """

    def __init__(self, dpi: int = 200):
        self.dpi = dpi

    def _backend(self):
        import PIL.Image  # noqa: F401
        import pymupdf  # noqa: F401
        import pytesseract

        return PIL.Image, pymupdf, pytesseract

    def available(self) -> bool:
        """True when Tesseract + rasterization libs are usable."""
        try:
            Pillow_Image, pymupdf, pytesseract = self._backend()
            pytesseract.get_tesseract_version()
            pymupdf.__version__
            Pillow_Image.__name__
            return True
        except Exception:
            return False

    def pdf_to_text(self, data: bytes) -> str:
        """Read all pages of a PDF via OCR and return the combined text."""
        PIL_Image, pymupdf, pytesseract = self._backend()
        document = pymupdf.open(stream=data, filetype="pdf")
        parts = []
        for page in document:
            pix = page.get_pixmap(dpi=self.dpi)
            image = PIL_Image.open(io.BytesIO(pix.tobytes("png"))).convert("L")
            parts.append(self._read_image(pytesseract, image))
        return "\n".join(parts).strip()

    @staticmethod
    def _read_image(pytesseract, image) -> str:
        """Tesseract OCR with German+English hint, falling back to defaults."""
        try:
            return pytesseract.image_to_string(image, lang="deu+eng")
        except Exception:
            return pytesseract.image_to_string(image)


def extract_cv_text(data: bytes,
                    filename: str,
                    force_ocr: bool = False,
                    ocr_engine=None) -> CVTextResult:
    """Convert raw document bytes into plain text.

    Uses fast extraction first; runs OCR only when the fast result is
    clearly insufficient (or ``force_ocr`` is set), so processing time is
    not wasted on documents that extract well. ``ocr_engine`` allows tests
    to inject a fake OCR implementation (callable ``data -> str`` or an
    object with ``pdf_to_text``).
    """
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return _extract_pdf(data, force_ocr=force_ocr, ocr_engine=ocr_engine)
    if name.endswith(".docx"):
        return _extract_docx(data)
    if name.endswith(".txt"):
        text = data.decode("utf-8", errors="replace").strip()
        return CVTextResult(text, "txt")
    return CVTextResult("", "unsupported")


def _extract_pdf(data: bytes, force_ocr: bool = False, ocr_engine=None) -> CVTextResult:
    if PdfReader is None:
        return CVTextResult("", "pdf_pkg_missing")

    try:
        reader = PdfReader(io.BytesIO(data))
        text_parts = []
        for page in reader.pages:
            try:
                text_parts.append(page.extract_text() or "")
            except Exception:
                text_parts.append("")
        text = "\n".join(text_parts).strip()
    except Exception:
        return CVTextResult("", "pdf_empty")

    if force_ocr or is_insufficient(text):
        engine = ocr_engine or OCREngine()
        if engine is not None and _engine_available(engine):
            try:
                if callable(engine):
                    ocr_text = engine(data)
                else:
                    ocr_text = engine.pdf_to_text(data)
            except Exception:
                ocr_text = ""
            ocr_text = (ocr_text or "").strip()
            if ocr_text and not is_insufficient(ocr_text):
                return CVTextResult(ocr_text, "pdf_ocr")
            # OCR did not help: return whatever fast extraction found.
            return CVTextResult(text, "pdf_text")
        if not text:
            return CVTextResult("", "pdf_ocr_missing")
        return CVTextResult(text, "pdf_text")

    return CVTextResult(text, "pdf_text")


def _engine_available(engine) -> bool:
    if callable(engine):
        return True
    try:
        return bool(engine.available())
    except Exception:
        return False


def _extract_docx(data: bytes) -> CVTextResult:
    if DocxDocument is None:
        return CVTextResult("", "docx_pkg_missing")
    try:
        document = DocxDocument(io.BytesIO(data))
        text_parts = []
        for para in document.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    text_parts.append(" | ".join(cells))
        return CVTextResult("\n".join(text_parts).strip(), "docx")
    except Exception:
        return CVTextResult("", "docx")